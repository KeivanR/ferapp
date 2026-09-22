"""Convertit la table Ciqual (xlsx) en foods.csv léger, prêt à être chargé par l'app.

Usage :  python build_foods.py "Table Ciqual 2025_FR_2025_11_03.xlsx" [foods.csv]
Nécessite : pip install openpyxl   (uniquement sur ton ordinateur, pas dans l'app)

Tout le paramétrage est dans config.toml :
  - [nutrients.<clé>] csv_column et ciqual : quelle colonne Ciqual alimente quelle colonne du CSV
  - [foods_build] : aliments écartés, traitement de "traces" et de "< x"

Règles de nettoyage des valeurs Ciqual (par 100 g) :
  - "-" ou cellule vide  -> valeur manquante (case vide dans le CSV, comptée 0 par l'app)
  - "traces"             -> traces_value (0 par défaut)
  - "< 0,25"             -> 0,25 x below_limit_factor (0 par défaut : on ne surestime pas les apports)
  - "2,45"               -> 2.45 (virgule décimale française)
"""

from __future__ import annotations

import csv
import re
import sys
import warnings
from pathlib import Path

import openpyxl

from config import load_config


def parse_value(raw, below_limit_factor: float = 0.0, traces_value: float = 0.0) -> float | None:
    """Valeur Ciqual -> float, ou None si manquante."""
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    text = " ".join(str(raw).split()).lower()
    if text in ("", "-"):
        return None
    if text == "traces":
        return traces_value
    if text.startswith("<"):
        try:
            limit = float(text[1:].strip().replace(",", "."))
        except ValueError:
            return 0.0
        return limit * below_limit_factor
    return float(text.replace(",", "."))


def resolve(row, alternatives: list[list[int]], **parse_options) -> float | None:
    """Première alternative ayant une valeur ; plusieurs colonnes dans une alternative = somme."""
    for cols in alternatives:
        parts = [v for v in (parse_value(row[c], **parse_options) for c in cols) if v is not None]
        if parts:
            return sum(parts)
    return None


def clean_header(h) -> str:
    return " ".join(str(h or "").split())


def convert(xlsx_path: str | Path, csv_path: str | Path, config: dict | None = None) -> dict:
    config = config or load_config()
    nutrients = config["nutrients"]
    build = config["foods_build"]
    options = {
        "below_limit_factor": build["below_limit_factor"],
        "traces_value": build["traces_value"],
    }

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # avertissement inoffensif sur l'en-tête/pied de page
        wb = openpyxl.load_workbook(xlsx_path, read_only=True)
        rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    headers = [clean_header(h) for h in rows[0]]

    def find_col(pattern: str, nutrient: str) -> int:
        found = [i for i, h in enumerate(headers) if re.search(pattern, h)]
        if len(found) != 1:
            raise SystemExit(
                f"[nutrients.{nutrient}] ciqual : {len(found)} colonnes pour {pattern!r} (attendu : 1)"
            )
        return found[0]

    # nutriment -> liste d'alternatives, chacune = liste d'index de colonnes (plusieurs = somme)
    idx = {}
    for n in nutrients:
        if not n["ciqual"]:
            raise SystemExit(f"[nutrients.{n['key']}] : « ciqual » est vide, impossible de le convertir")
        idx[n["key"]] = [
            [find_col(p, n["key"]) for p in ((alt,) if isinstance(alt, str) else alt)]
            for alt in n["ciqual"]
        ]
    name_i = headers.index("alim_nom_fr")
    code_i = headers.index("alim_code")

    # Un aliment est inutile si aucun nutriment du groupe requis n'est renseigné.
    required = [i for i, n in enumerate(nutrients) if n["group"] == build["required_group"]]

    kept = dropped = 0
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["alim_code", "aliment", *(n["col"] for n in nutrients)])
        for r in rows[1:]:
            name = (r[name_i] or "").strip()
            values = [resolve(r, idx[n["key"]], **options) for n in nutrients]
            if not name or (required and all(values[i] is None for i in required)):
                dropped += 1
                continue
            w.writerow([r[code_i], name, *("" if v is None else f"{v:g}" for v in values)])
            kept += 1
    return {"kept": kept, "dropped": dropped}


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    config = load_config()
    out = sys.argv[2] if len(sys.argv) > 2 else Path(__file__).parent / config["app"]["foods_file"]
    print(convert(sys.argv[1], out, config), "->", out)
